import json
from datetime import date, timedelta
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree
import re

from lerntagebuch import config

# Namespaces for XML processing
NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'w14': 'http://schemas.microsoft.com/office/word/2010/wordml',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}

def get_lernort_pattern(data, default=None):
    """Determines the 5-character Lernort pattern."""
    pattern = data.get("lernort")
    if pattern:
        pattern = pattern.strip("()")
        if len(pattern) == 5:
            return pattern.upper()
    return default or config.DEFAULT_PATTERN

def calculate_hours(pattern, hours_table=None):
    """Stunden je Tag als Text fuer das Dokument, z. B. ["10", "8,5", ...].

    hours_table ordnet jedem Kuerzel fuenf Werte zu (Mo-Fr), siehe
    config.DEFAULT_HOURS. Unbekannte Kuerzel ergeben "0".
    """
    table = hours_table or config.DEFAULT_HOURS
    hours = []
    for i, char in enumerate(pattern):
        values = table.get(char)
        hours.append(config.format_hours(values[i]) if values else "0")
    return hours

def replace_text_in_paragraph(p, replacements):
    """Replaces text placeholders in a paragraph, attempting to preserve runs."""
    t_xpath = etree.XPath('.//w:t', namespaces=NS)
    all_t = t_xpath(p._element)
    if not all_t:
        return

    # 1. Surgical replacement (preserves formatting if placeholder is in one tag)
    for t in all_t:
        if t.text:
            for target, text in replacements.items():
                if target in t.text:
                    t.text = t.text.replace(target, text)

    # 2. Robust replacement (handles split placeholders by merging runs ONLY in this paragraph)
    full_text = "".join(t.text for t in all_t if t.text)
    needs_merge = False
    for target in replacements.keys():
        if target in full_text:
            needs_merge = True
            break
            
    if needs_merge:
        for target, text in replacements.items():
            full_text = full_text.replace(target, text)
        all_t[0].text = full_text
        for t in all_t[1:]:
            t.text = ""

def checkbox_label(sdt):
    """Beschriftung eines Kontrollkaestchens: Text dahinter bis zum naechsten Kaestchen."""
    parts = []
    for sib in sdt.itersiblings():
        if sib.tag == '{%s}sdt' % NS['w']:
            break
        parts += [t.text or "" for t in sib.iter('{%s}t' % NS['w'])]
    return "".join(parts).strip()

def process_checkboxes(cell, target_char):
    """Sets the correct checkbox in a 'Lernort' cell based on the pattern character.

    Die Kaestchen werden ueber ihre Beschriftung zugeordnet
    (config.LERNORT_CHECKBOX_LABELS), die Reihenfolge im Template ist egal.
    """
    labels = config.LERNORT_CHECKBOX_LABELS
    target = labels.get(target_char)
    sdt_xpath = etree.XPath('.//w:sdt', namespaces=NS)
    cb_xpath = etree.XPath('.//w14:checkbox', namespaces=NS)
    checked_xpath = etree.XPath('.//w14:checked', namespaces=NS)
    t_xpath = etree.XPath('.//w:sdtContent//w:t', namespaces=NS)
    sdts = sdt_xpath(cell._element)
    for sdt in sdts:
        if not cb_xpath(sdt):
            continue
        label = checkbox_label(sdt)
        if not any(label.startswith(l) for l in labels.values()):
            continue
        is_checked = target is not None and label.startswith(target)
        checked_elems = checked_xpath(sdt)
        if not checked_elems:
            cb = cb_xpath(sdt)[0]
            checked_elem = etree.SubElement(cb, '{%s}checked' % NS['w14'])
            checked_elem.set('{%s}val' % NS['w14'], '1' if is_checked else '0')
        else:
            for checked_elem in checked_elems:
                checked_elem.set('{%s}val' % NS['w14'], '1' if is_checked else '0')
        for t_elem in t_xpath(sdt):
            t_elem.text = '☒' if is_checked else '☐'

def insert_signature(doc, image_path, text=""):
    """Setzt das Unterschriftsbild in den Bild-Platzhalter des Templates.

    Der Platzhalter ist ein transparentes Bild mit dem Alternativtext
    {{UNTERSCHRIFT}}. Rahmen, Position und Textfluss bleiben erhalten, nur der
    Bildinhalt wird getauscht und mit Seitenverhaeltnis in den Rahmen
    eingepasst. Ohne Bild bleibt der transparente Platzhalter stehen, der Platz
    ist dann genauso reserviert.

    Mit text (Nachname als Unterschrift) steht der Text direkt vor dem
    Platzhalter, ein Bild wird dann nicht eingesetzt.
    """
    embed = f"{{{NS['r']}}}embed"
    w = f"{{{NS['w']}}}"
    frames = [pr.getparent() for pr in doc.element.body.iter(f"{{{NS['wp']}}}docPr")
              if pr.get("descr") == config.SIGNATURE_PLACEHOLDER]

    for frame in frames:  # wp:inline oder wp:anchor
        if text:
            drawing_run = next(a for a in frame.iterancestors() if a.tag == w + "r")
            run = etree.Element(w + "r")
            etree.SubElement(run, w + "t").text = text
            drawing_run.addprevious(run)
        elif image_path:
            blip = frame.find(f".//{{{NS['a']}}}blip")
            old_rid = blip.get(embed)
            rid, image = doc.part.get_or_add_image(str(image_path))
            blip.set(embed, rid)
            if old_rid != rid:
                doc.part.drop_rel(old_rid)

            # Zuschnitt gehoert zum alten Bild, nicht zur neuen Unterschrift
            for src_rect in list(frame.iter(f"{{{NS['a']}}}srcRect")):
                src_rect.getparent().remove(src_rect)

            # In den bisherigen Rahmen einpassen, Seitenverhaeltnis beibehalten
            extent = frame.find(f"{{{NS['wp']}}}extent")
            box_cx, box_cy = int(extent.get("cx")), int(extent.get("cy"))
            scale = min(box_cx / image.px_width, box_cy / image.px_height)
            cx = str(round(image.px_width * scale))
            cy = str(round(image.px_height * scale))
            extent.set("cx", cx)
            extent.set("cy", cy)
            for ext in frame.iterfind(f".//{{{NS['a']}}}xfrm/{{{NS['a']}}}ext"):
                ext.set("cx", cx)
                ext.set("cy", cy)

        # Alternativtext im fertigen Dokument neutral halten
        for pr in frame.iter(f"{{{NS['wp']}}}docPr", f"{{{NS['pic']}}}cNvPr"):
            pr.set("descr", "Unterschrift")


def output_path(output_dir, kw, year):
    """Pfad des erzeugten Dokuments, z. B. output/Lerntagebuch_KW39_2026.docx."""
    return Path(output_dir) / f"Lerntagebuch_KW{int(kw)}_{int(year)}.docx"

def generate_document(week_path=None, template_path=None, output_dir=None,
                      settings=None):
    """Erzeugt das Dokument und gibt den Pfad zurueck (None bei Fehlern).

    Ohne Angaben gelten die Pfade aus config.py und die Einstellungen aus
    settings.json. Die GUI uebergibt ihre bereits geladenen Einstellungen, die
    Tests ihre eigenen, damit settings.json des Nutzers keine Rolle spielt.
    """
    if settings is None:
        settings, warnings = config.load_settings()
        for w in warnings:
            print(f"Warnung: {w}")

    if output_dir is None:
        output_dir, fallback = config.resolve_output_dir(settings)
        if fallback:
            print(f"Hinweis: Speicherort '{settings['output_dir']}' existiert nicht, "
                  f"gespeichert wird in {output_dir}.")

    week_path = Path(week_path or config.WEEK_PATH)
    template_path = Path(template_path or config.TEMPLATE_PATH)
    output_dir = Path(output_dir)

    try:
        with open(week_path, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except FileNotFoundError:
        print("Error: week.json not found.")
        return None

    KW = int(data.get("kw", 1))
    YEAR = int(data.get("jahr", 2026))
    monday = date.fromisocalendar(YEAR, KW, 1)
    signature_text = config.signature_text(settings)
    signature_path, signature_missing = config.resolve_signature(settings)
    if signature_text:
        signature_path, signature_missing = None, False  # Name statt Bild
    if signature_missing:
        print(f"Hinweis: Unterschriftsbild '{settings['signature']}' nicht gefunden, "
              "das Unterschriftsfeld bleibt leer.")

    pattern = get_lernort_pattern(data, settings["default_pattern"])
    h_list = calculate_hours(pattern, settings["hours"])

    lf_raw = data.get("lf", "").strip()
    lf_base = ""
    if lf_raw:
        # Bei mehreren Lernfeld-Nummern automatisch den Plural verwenden.
        word = "Lernfelder" if len(re.findall(r'\d+', lf_raw)) > 1 else "Lernfeld"
        if lf_raw.upper().startswith("LF"):
            lf_base = re.sub(r'^LF\s*', word + ' ', lf_raw, flags=re.IGNORECASE)
        elif lf_raw.upper().startswith("LERNFELD"):
            lf_base = lf_raw
        else:
            lf_base = f"{word} {lf_raw}"

    days_key = ["MONTAG", "DIENSTAG", "MITTWOCH", "DONNERSTAG", "FREITAG"]
    replacements = {
        "{{KW}}": str(KW),
        "{{NAME}}": config.full_name(settings),
        "{{DATUM_MO}}": monday.strftime("%d.%m.%Y"),
        "{{DATUM_DI}}": (monday + timedelta(days=1)).strftime("%d.%m.%Y"),
        "{{DATUM_MI}}": (monday + timedelta(days=2)).strftime("%d.%m.%Y"),
        "{{DATUM_DO}}": (monday + timedelta(days=3)).strftime("%d.%m.%Y"),
        "{{DATUM_FR}}": (monday + timedelta(days=4)).strftime("%d.%m.%Y"),
        "{{H_MO}}": h_list[0],
        "{{H_DI}}": h_list[1],
        "{{H_MI}}": h_list[2],
        "{{H_DO}}": h_list[3],
        "{{H_FR}}": h_list[4],
    }
    # Bei Feiertag (F) oder Urlaub (U) den Tagesinhalt ignorieren und
    # stattdessen die Abwesenheitsart als Inhalt eintragen.
    off_day_labels = {'F': "Feiertag", 'U': "Urlaub"}
    for i, day in enumerate(days_key):
        if pattern[i] in off_day_labels:
            content = off_day_labels[pattern[i]]
        else:
            content = data.get(day.lower(), "")
        replacements[f"{{{{{day}}}}}"] = content

        placeholder = f"{{{{LF_{day}}}}}"
        replacements[placeholder] = lf_base if (lf_base and pattern[i] in ('S', 'M')) else ""

    try:
        doc = Document(str(template_path))
    except Exception as e:
        print(f"Error loading template.docx: {e}")
        return None

    # 1. Process Header
    for section in doc.sections:
        header = section.header
        # Paragraphs in header
        for p in header.paragraphs:
            replace_text_in_paragraph(p, replacements)
            if str(KW) in p.text and "{{KW}}" not in p.text:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(16)
        # Tables in header
        for table in header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        replace_text_in_paragraph(p, replacements)

    # 2. Process Tables (Body)
    lernort_count = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_text_in_paragraph(p, replacements)
                if "Lernort:" in cell.text and lernort_count < len(pattern):
                    process_checkboxes(cell, pattern[lernort_count])
                    lernort_count += 1

    # 3. Process Body Paragraphs
    for p in doc.paragraphs:
        replace_text_in_paragraph(p, replacements)

    # 4. Unterschrift
    insert_signature(doc, signature_path, signature_text)

    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_path(output_dir, KW, YEAR)
    doc.save(str(out_path))
    print(f"Successfully generated: {out_path}")
    return str(out_path)

if __name__ == "__main__":
    generate_document()
