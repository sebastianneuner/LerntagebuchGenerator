"""
Tests fuer die bestehende Logik des Generators (lerntagebuch/generator.py).

Sie halten das heutige Verhalten fest, damit beim Umbau sofort auffaellt,
wenn sich etwas ungewollt aendert.

Ausfuehren (im Projektordner):
    .venv\\Scripts\\python.exe -m unittest discover -s tests   (Windows)
    .venv/bin/python -m unittest discover -s tests             (Linux/macOS)
"""

import io
import json
import re
import shutil
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

from docx import Document

from helpers import make_png
from lerntagebuch import config, generator

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"

# Tabellen 0-4 im Dokument sind Montag bis Freitag.
DAY_TABLES = range(0, 5)


class TestCalculateHours(unittest.TestCase):

    def test_schule_und_mobil_mo_bis_do_10_freitag_7(self):
        self.assertEqual(generator.calculate_hours("SSSSS"), ["10", "10", "10", "10", "7"])
        self.assertEqual(generator.calculate_hours("MMMMM"), ["10", "10", "10", "10", "7"])

    def test_betrieb_mo_bis_do_8_5_freitag_6(self):
        self.assertEqual(generator.calculate_hours("BBBBB"), ["8,5", "8,5", "8,5", "8,5", "6"])

    def test_feiertag_und_urlaub_0(self):
        self.assertEqual(generator.calculate_hours("FFFFF"), ["0"] * 5)
        self.assertEqual(generator.calculate_hours("UUUUU"), ["0"] * 5)

    def test_unbekanntes_zeichen_0(self):
        self.assertEqual(generator.calculate_hours("XXXXX"), ["0"] * 5)

    def test_standardmuster(self):
        self.assertEqual(generator.calculate_hours("MSSSM"), ["10", "10", "10", "10", "7"])

    def test_gemischt(self):
        self.assertEqual(generator.calculate_hours("SBMFU"), ["10", "8,5", "10", "0", "0"])

    def test_eigene_tabelle_je_tag(self):
        tabelle = dict(config.DEFAULT_HOURS, B=[8, 7.5, 8, 8, 4.25], U=[8, 8, 8, 8, 8])
        self.assertEqual(generator.calculate_hours("BBUBB", tabelle),
                         ["8", "7,5", "8", "8", "4,25"])


class TestGetLernortPattern(unittest.TestCase):

    def test_fehlt_ergibt_standard(self):
        self.assertEqual(generator.get_lernort_pattern({}), "MSSSM")

    def test_leer_ergibt_standard(self):
        self.assertEqual(generator.get_lernort_pattern({"lernort": ""}), "MSSSM")

    def test_kleinbuchstaben_werden_gross(self):
        self.assertEqual(generator.get_lernort_pattern({"lernort": "sbmfu"}), "SBMFU")

    def test_klammern_werden_entfernt(self):
        self.assertEqual(generator.get_lernort_pattern({"lernort": "(BSSSM)"}), "BSSSM")

    def test_falsche_laenge_ergibt_standard(self):
        self.assertEqual(generator.get_lernort_pattern({"lernort": "SSS"}), "MSSSM")
        self.assertEqual(generator.get_lernort_pattern({"lernort": "SSSSSS"}), "MSSSM")


class TestGenerateDocument(unittest.TestCase):
    """Erzeugt echte Dokumente in einem Temp-Ordner und prueft den Inhalt."""

    def setUp(self):
        self._tmp = Path(tempfile.mkdtemp())
        self.week = self._tmp / "week.json"
        self.out_dir = self._tmp / "output"
        # Eigene Einstellungen statt settings.json des Nutzers
        self.settings = config.default_settings()

    def tearDown(self):
        shutil.rmtree(self._tmp)

    def run_generator(self):
        with redirect_stdout(io.StringIO()):
            return generator.generate_document(
                week_path=self.week,
                template_path=config.TEMPLATE_PATH,
                output_dir=self.out_dir,
                settings=self.settings)

    # ------------------------------------------------------------ Helfer
    def generate(self, **overrides):
        """Schreibt week.json, erzeugt das Dokument und gibt es zurueck."""
        data = {
            "jahr": 2026, "kw": 38, "lernort": "SBMFU", "lf": "LF 3",
            "montag": "Inhalt Mo", "dienstag": "Inhalt Di",
            "mittwoch": "Inhalt Mi", "donnerstag": "Inhalt Do",
            "freitag": "Inhalt Fr",
        }
        data.update(overrides)
        with open(self.week, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        out = self.run_generator()
        self.assertIsNotNone(out)
        return Document(out)

    @staticmethod
    def day(doc, i):
        """Liefert Stunden, Inhalt, Lernfeld und Lernort-Haekchen eines Tages."""
        t = doc.tables[DAY_TABLES[i]]
        lernort_cell = t.cell(0, 1)
        # Haekchen in fester Reihenfolge S, B, M – ueber die Beschriftung,
        # nicht ueber die Position im Template
        by_label = {}
        for sdt in lernort_cell._element.iter(f"{{{W}}}sdt"):
            checked = next(sdt.iter(f"{{{W14}}}checked"), None)
            if checked is not None:
                by_label[generator.checkbox_label(sdt)] = checked.get(f"{{{W14}}}val")
        checks = [next(v for k, v in by_label.items() if k.startswith(label))
                  for label in config.LERNORT_CHECKBOX_LABELS.values()]
        return {
            "hours": t.cell(1, 2).text.strip(),
            "content": t.cell(2, 1).text,
            "lf": t.cell(4, 1).text,
            "lernort": checks,
        }

    @staticmethod
    def all_text(doc):
        """Gesamter Text aus Body-Tabellen, Absaetzen und Kopfzeile."""
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            parts += [c.text for row in t.rows for c in row.cells]
        for s in doc.sections:
            parts += [p.text for p in s.header.paragraphs]
            for t in s.header.tables:
                parts += [c.text for row in t.rows for c in row.cells]
        return "\n".join(parts)

    # ------------------------------------------------------------- Tests
    def test_keine_platzhalter_uebrig(self):
        doc = self.generate()
        self.assertEqual(re.findall(r"\{\{[^}]*\}\}", self.all_text(doc)), [])

    def test_name_in_kopfzeile(self):
        doc = self.generate()
        kopf = [c.text for s in doc.sections for t in s.header.tables
                for row in t.rows for c in row.cells]
        self.assertIn(f"Name  {config.DEFAULT_LAST_NAME}, {config.DEFAULT_FIRST_NAME}", kopf)

    def test_dateiname(self):
        self.generate()
        self.assertTrue((self.out_dir / "Lerntagebuch_KW38_2026.docx").is_file())
        # Die GUI prueft vorab mit output_path, ob die Woche schon existiert
        self.assertTrue(generator.output_path(self.out_dir, "38", "2026").is_file())

    def test_kw_und_datum(self):
        doc = self.generate()
        kopf = doc.sections[0].header.tables[0]
        self.assertEqual(kopf.cell(0, 1).text, "KW 38")
        self.assertEqual(kopf.cell(1, 1).text,
                         "Unterrichtswoche  14.09.2026 – 18.09.2026")

    def test_stunden(self):
        doc = self.generate(lernort="SBMFU")
        self.assertEqual([self.day(doc, i)["hours"] for i in range(5)],
                         ["10", "8,5", "10", "0", "0"])

    def test_stunden_freitag_sonderfall(self):
        doc = self.generate(lernort="MSSSM")
        self.assertEqual(self.day(doc, 4)["hours"], "7")

    def test_inhalt_und_feiertag_urlaub(self):
        doc = self.generate(lernort="SBMFU")
        self.assertEqual([self.day(doc, i)["content"] for i in range(5)],
                         ["Inhalt Mo", "Inhalt Di", "Inhalt Mi", "Feiertag", "Urlaub"])

    def test_lernfeld_nur_an_schule_und_mobil(self):
        doc = self.generate(lernort="SBMFU", lf="LF 3")
        self.assertEqual([self.day(doc, i)["lf"] for i in range(5)],
                         ["Lernfeld 3", "", "Lernfeld 3", "", ""])

    def test_lernfeld_schreibweisen(self):
        faelle = {
            "LF 3": "Lernfeld 3",
            "lf3": "Lernfeld 3",
            "3": "Lernfeld 3",
            "3, 4": "Lernfelder 3, 4",
            "LF 3/4": "Lernfelder 3/4",
            "Lernfeld 5": "Lernfeld 5",
            "": "",
        }
        for eingabe, erwartet in faelle.items():
            with self.subTest(eingabe=eingabe):
                doc = self.generate(lernort="SSSSS", lf=eingabe)
                self.assertEqual(self.day(doc, 0)["lf"], erwartet)

    def test_lernort_checkboxen(self):
        doc = self.generate(lernort="SBMFU")
        self.assertEqual([self.day(doc, i)["lernort"] for i in range(5)], [
            ["1", "0", "0"],   # S = Akademie/Standort
            ["0", "1", "0"],   # B = Betrieb
            ["0", "0", "1"],   # M = mobil
            ["0", "0", "0"],   # F
            ["0", "0", "0"],   # U
        ])

    def test_vermittelt_durch_bleibt_unangetastet(self):
        doc = self.generate()
        for i in range(5):
            with self.subTest(tag=i):
                zelle = doc.tables[DAY_TABLES[i]].cell(5, 1)
                checks = [x.get(f"{{{W14}}}val")
                          for x in zelle._element.iter(f"{{{W14}}}checked")]
                self.assertTrue(checks)
                self.assertEqual(set(checks), {"0"})

    def test_umlaute_bleiben_erhalten(self):
        doc = self.generate(montag="Übung für Größenänderung – ß")
        self.assertEqual(self.day(doc, 0)["content"], "Übung für Größenänderung – ß")

    def test_week_json_mit_bom(self):
        # Windows-Editoren speichern UTF-8 manchmal mit BOM.
        data = {"jahr": 2026, "kw": 38, "lernort": "SSSSS", "montag": "Mit BOM"}
        with open(self.week, "w", encoding="utf-8-sig") as f:
            json.dump(data, f)
        out = self.run_generator()
        self.assertEqual(self.day(Document(out), 0)["content"], "Mit BOM")

    def test_ohne_week_json(self):
        self.assertIsNone(self.run_generator())

    # --------------------------------------------- mit eigenen Einstellungen
    def test_eigener_name(self):
        self.settings.update(last_name="Mustermann", first_name="Max")
        doc = self.generate()
        kopf = [c.text for s in doc.sections for t in s.header.tables
                for row in t.rows for c in row.cells]
        self.assertIn("Name  Mustermann, Max", kopf)

    def test_eigene_stunden(self):
        self.settings["hours"]["S"] = [6, 6, 6, 6, 3.5]
        doc = self.generate(lernort="SSSSS")
        self.assertEqual([self.day(doc, i)["hours"] for i in range(5)],
                         ["6", "6", "6", "6", "3,5"])

    def test_standard_lernort_aus_einstellungen(self):
        # week.json ohne gueltigen Lernort -> Standard aus den Einstellungen
        self.settings["default_pattern"] = "BBBBB"
        doc = self.generate(lernort="")
        self.assertEqual(self.day(doc, 0)["lernort"], ["0", "1", "0"])

    def test_speicherort_aus_einstellungen(self):
        eigener = self._tmp / "Abgaben"
        eigener.mkdir()
        self.settings["output_dir"] = str(eigener)
        data = {"jahr": 2026, "kw": 38, "lernort": "SSSSS", "montag": "x"}
        with open(self.week, "w", encoding="utf-8") as f:
            json.dump(data, f)
        with redirect_stdout(io.StringIO()):
            out = generator.generate_document(
                week_path=self.week, template_path=config.TEMPLATE_PATH,
                settings=self.settings)  # output_dir=None -> aus Einstellungen
        self.assertEqual(Path(out).parent, eigener)

    # ------------------------------------------------------- Unterschrift
    # Rahmen des Platzhalters im Template (EMU), siehe resources/template.docx
    BOX = (549680, 323850)

    @staticmethod
    def signature(doc):
        """Rahmen, Bilddaten und Zuschnitt des Unterschriftsbilds im Dokument."""
        NS = generator.NS
        pr = [p for p in doc.element.body.iter(f"{{{NS['wp']}}}docPr")
              if p.get("name") == "Unterschrift"]
        assert len(pr) == 1, "Unterschrift-Bild nicht gefunden"
        frame = pr[0].getparent()
        extent = frame.find(f"{{{NS['wp']}}}extent")
        xfrm = frame.find(f".//{{{NS['a']}}}xfrm/{{{NS['a']}}}ext")
        rid = frame.find(f".//{{{NS['a']}}}blip").get(f"{{{NS['r']}}}embed")
        return {
            "extent": (int(extent.get("cx")), int(extent.get("cy"))),
            "xfrm": (int(xfrm.get("cx")), int(xfrm.get("cy"))),
            "blob": doc.part.related_parts[rid].blob,
            "crop": frame.findall(f".//{{{NS['a']}}}srcRect"),
            "descr": pr[0].get("descr"),
            "images": [r for r in doc.part.rels.values() if "image" in r.reltype],
        }

    def with_signature(self, width, height):
        png = make_png(width, height)
        path = self._tmp / "sig.png"
        path.write_bytes(png)
        self.settings["signature"] = str(path)
        return png

    def test_ohne_unterschrift_bleibt_platzhalter(self):
        sig = self.signature(self.generate())
        self.assertEqual(sig["extent"], self.BOX)
        self.assertLess(len(sig["blob"]), 1000)  # transparenter Platzhalter
        self.assertEqual(sig["descr"], "Unterschrift")

    def test_breite_unterschrift_passt_in_rahmen(self):
        png = self.with_signature(300, 100)  # 3:1, breiter als der Rahmen
        sig = self.signature(self.generate())
        self.assertEqual(sig["blob"], png)
        cx, cy = sig["extent"]
        self.assertEqual(cx, self.BOX[0])          # volle Breite
        self.assertLess(cy, self.BOX[1])            # dafuer niedriger
        self.assertAlmostEqual(cx / cy, 3, places=2)
        self.assertEqual(sig["xfrm"], sig["extent"])
        self.assertEqual(sig["crop"], [])
        self.assertEqual(len(sig["images"]), 1)     # Platzhalter entfernt

    def test_hohe_unterschrift_passt_in_rahmen(self):
        self.with_signature(100, 300)  # 1:3
        cx, cy = self.signature(self.generate())["extent"]
        self.assertEqual(cy, self.BOX[1])           # volle Hoehe
        self.assertLess(cx, self.BOX[0])
        self.assertAlmostEqual(cy / cx, 3, places=2)

    def test_unterschrift_fehlt_dokument_trotzdem(self):
        self.settings["signature"] = str(self._tmp / "gibt_es_nicht.png")
        sig = self.signature(self.generate())
        self.assertEqual(sig["extent"], self.BOX)
        self.assertLess(len(sig["blob"]), 1000)

    def signature_paragraph_text(self, doc):
        """Text des Absatzes mit dem Unterschrift-Platzhalter."""
        pr = next(p for p in doc.element.body.iter(f"{{{generator.NS['wp']}}}docPr")
                  if p.get("name") == "Unterschrift")
        para = next(a for a in pr.iterancestors() if a.tag == f"{{{W}}}p")
        return "".join(t.text or "" for t in para.iter(f"{{{W}}}t"))

    def test_nachname_als_unterschrift(self):
        self.settings.update(last_name="Mustermann", first_name="Max")
        self.settings["signature_name"] = True
        doc = self.generate()
        self.assertEqual(self.signature_paragraph_text(doc).split(),
                         ["18.09.2026,", "Mustermann"])
        # Platzhalter bleibt stehen, das Layout aendert sich nicht
        self.assertEqual(self.signature(doc)["extent"], self.BOX)

    def test_nachname_statt_bild(self):
        self.with_signature(300, 100)
        self.settings.update(last_name="Mustermann", first_name="Max")
        self.settings["signature_name"] = True
        doc = self.generate()
        sig = self.signature(doc)
        self.assertLess(len(sig["blob"]), 1000)  # kein Bild, nur der Platzhalter
        self.assertIn("Mustermann", self.signature_paragraph_text(doc))

    def test_ohne_nachname_option_kein_text(self):
        self.settings.update(last_name="Mustermann", first_name="Max")
        doc = self.generate()
        self.assertNotIn("Mustermann", self.signature_paragraph_text(doc))

    def test_unterschrift_aendert_restliches_dokument_nicht(self):
        self.with_signature(300, 100)
        doc = self.generate()
        self.assertEqual(re.findall(r"\{\{[^}]*\}\}", self.all_text(doc)), [])
        self.assertEqual([self.day(doc, i)["hours"] for i in range(5)],
                         ["10", "8,5", "10", "0", "0"])


if __name__ == "__main__":
    unittest.main()
